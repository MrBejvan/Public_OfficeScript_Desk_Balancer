from datetime import date

from openpyxl import load_workbook

from core.exporter import export_schedule
from core.models import Desk, Employee, ShiftRequest
from core.scheduler import build_schedule


def test_export_produces_expected_sheets(tmp_path):
    start = date(2026, 7, 1)
    end = date(2026, 7, 31)
    requests = [
        ShiftRequest(Employee("Alice"), date(2026, 7, 1), Desk.EAST),
        ShiftRequest(Employee("Bob"), date(2026, 7, 1), Desk.WEST),
    ]
    period = build_schedule(requests, start, end)

    out_path = tmp_path / "calendar.xlsx"
    export_schedule(period, out_path)

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert "Jul 2026" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Needs Attention" in wb.sheetnames  # shortfalls exist by design here


def test_export_spans_multiple_month_sheets(tmp_path):
    start = date(2026, 6, 25)
    end = date(2026, 7, 5)
    period = build_schedule([], start, end)

    out_path = tmp_path / "calendar.xlsx"
    export_schedule(period, out_path)

    wb = load_workbook(out_path)
    assert "Jun 2026" in wb.sheetnames
    assert "Jul 2026" in wb.sheetnames
