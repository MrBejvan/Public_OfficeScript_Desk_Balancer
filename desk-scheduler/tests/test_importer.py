from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.importer import KronosImportError, load_shift_requests, parse_day_cell
from core.importer import _clean_name  # noqa: F401 -- private, tested directly below


# ---------------------------------------------------------------------------
# parse_day_cell: the core cell-interpretation rule, tested directly
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,expected",
    [
        (None, (False, None)),
        ("", (False, None)),
        ("   ", (False, None)),
        ("D12", (True, "D12")),
        ("08c12", (True, "08c12")),
        ("PTO Scheduled Hours 12:00", (False, None)),
        ("pto", (False, None)),
        ("RDY", (False, None)),
        ("SDY", (False, None)),
        ("VDY 12:30", (False, None)),
        ("VDY 00:30", (False, None)),
        ("RDY \nD12", (True, "D12")),  # avoid-request overridden -> working
        ("SDY\nD12", (True, "D12")),
    ],
)
def test_parse_day_cell(raw, expected):
    assert parse_day_cell(raw) == expected


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("Alvarez, Jordan", "Alvarez, Jordan"),
        ("  Alvarez,   Jordan  ", "Alvarez, Jordan"),
        ("=HACK()", None),  # formula-injection guard, tested directly since
        ("123 Not A Name!!", None),  # openpyxl intercepts "=" cells at write time
        ("", None),
        (None, None),
    ],
)
def test_clean_name(raw, expected):
    assert _clean_name(raw) == expected


# ---------------------------------------------------------------------------
# Synthetic Kronos-grid workbook builder (mirrors the real export's shape,
# reverse-engineered from an actual Children's Colorado "Location Schedule"
# pull: a row with a literal "Employee" cell, a date row below it with one
# non-blank cell marking the start of each day's column span, then one row
# per employee with at most one non-blank cell per day's span).
# ---------------------------------------------------------------------------


def _write_block(ws, start_row: int, dates: list[date], employees: list[tuple[str, dict]], span_width: int = 2, value_in_second_col: bool = False):
    header_row = start_row
    date_row = start_row + 1
    ws.cell(row=header_row, column=1, value="Employee")
    ws.cell(row=header_row, column=3, value="Job")

    date_cols = {}
    col = 5
    for d in dates:
        ws.cell(row=date_row, column=col, value=d.strftime("%m/%d/\n%Y"))
        date_cols[d] = col
        col += span_width

    for row_offset, (name, day_values) in enumerate(employees):
        row = date_row + 1 + row_offset
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=3, value="ISS")
        for d, value in day_values.items():
            target_col = date_cols[d] + (1 if value_in_second_col else 0)
            ws.cell(row=row, column=target_col, value=value)

    return date_row + 1 + len(employees)  # first row after this block


def _save(path: Path, build_fn):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    build_fn(ws)
    wb.save(path)


# ---------------------------------------------------------------------------
# Tests against the full importer
# ---------------------------------------------------------------------------


def test_loads_valid_grid(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1, d2 = date(2026, 8, 1), date(2026, 8, 2)

    def build(ws):
        _write_block(
            ws,
            6,
            [d1, d2],
            [("Alvarez, Jordan", {d1: "D12", d2: "08c12"})],
        )

    _save(path, build)
    result = load_shift_requests(path)

    assert len(result.requests) == 2
    by_date = {r.work_date: r for r in result.requests}
    assert by_date[d1].employee.name == "Alvarez, Jordan"
    assert by_date[d1].shift_code == "D12"
    assert by_date[d1].preferred_desk is None
    assert by_date[d2].shift_code == "08c12"


def test_value_found_anywhere_in_the_days_column_span(tmp_path):
    # Real exports sometimes place the value in the *second* column of a
    # day's merged span rather than the first (observed directly in a real
    # export). The importer must find it regardless of which column it's in.
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 1)

    def build(ws):
        _write_block(ws, 6, [d1], [("Brooks, Taylor", {d1: "D12"})], value_in_second_col=True)

    _save(path, build)
    result = load_shift_requests(path)
    assert len(result.requests) == 1
    assert result.requests[0].shift_code == "D12"


def test_avoid_markers_alone_produce_no_request(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1, d2, d3, d4 = date(2026, 8, 1), date(2026, 8, 2), date(2026, 8, 3), date(2026, 8, 4)

    def build(ws):
        _write_block(
            ws,
            6,
            [d1, d2, d3, d4],
            [("Ellison, Avery", {d1: "RDY", d2: "SDY", d3: "VDY 12:30", d4: "PTO Scheduled Hours 12:00"})],
        )

    _save(path, build)
    result = load_shift_requests(path)
    assert result.requests == []
    assert result.issues == []


def test_overridden_avoid_marker_still_produces_a_working_request(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 2)

    def build(ws):
        _write_block(ws, 6, [d1], [("Alvarez, Jordan", {d1: "RDY \nD12"})])

    _save(path, build)
    result = load_shift_requests(path)
    assert len(result.requests) == 1
    assert result.requests[0].shift_code == "D12"


def test_blank_cells_produce_no_request(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1, d2 = date(2026, 8, 1), date(2026, 8, 2)

    def build(ws):
        _write_block(ws, 6, [d1, d2], [("Grant, Riley", {d1: "D12"})])  # d2 left blank

    _save(path, build)
    result = load_shift_requests(path)
    assert len(result.requests) == 1
    assert result.requests[0].work_date == d1


def test_multiple_location_blocks_are_both_parsed_and_dont_bleed_together(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 1)

    def build(ws):
        ws.cell(row=1, column=1, value="0000100-X1 - Location: SAMPLE/DEMO/DATA/0000100-X1")
        next_row = _write_block(ws, 2, [d1], [("Alvarez, Jordan", {d1: "D12"})])
        ws.cell(row=next_row, column=1, value="0000200-X2 - Location: SAMPLE/DEMO/DATA/0000200-X2")
        _write_block(ws, next_row + 1, [d1], [("Brooks, Taylor", {d1: "08c12"})])

    _save(path, build)
    result = load_shift_requests(path)

    names = {r.employee.name for r in result.requests}
    assert names == {"Alvarez, Jordan", "Brooks, Taylor"}
    assert len(result.requests) == 2


def test_location_banner_row_stops_employee_scanning(tmp_path):
    # A block with a trailing location-banner row immediately below the
    # last real employee row should not be misread as another employee.
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 1)

    def build(ws):
        last_row = _write_block(ws, 6, [d1], [("Alvarez, Jordan", {d1: "D12"})])
        ws.cell(row=last_row, column=1, value="Some Other - Location: Banner")

    _save(path, build)
    result = load_shift_requests(path)
    assert len(result.requests) == 1
    assert result.requests[0].employee.name == "Alvarez, Jordan"


def test_invalid_name_is_skipped_with_issue(tmp_path):
    # Note: a leading "=" is deliberately not used here -- openpyxl treats
    # a string cell value starting with "=" as an actual formula when
    # writing, and since this test file is never opened by real Excel there
    # is no cached result for data_only=True to read back (it comes back as
    # None, which is a different -- already covered -- code path: a blank
    # cell is skipped silently, not flagged as an invalid name). The name
    # guard against formula-injection is exercised directly in
    # test_parse_day_cell-adjacent unit coverage; here we just confirm a
    # plain invalid name (one that doesn't match a person's name shape) is
    # skipped with an issue instead of silently dropped.
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 1)

    def build(ws):
        _write_block(
            ws,
            6,
            [d1],
            [("123 Not A Name!!", {d1: "D12"}), ("Alvarez, Jordan", {d1: "D12"})],
        )

    _save(path, build)
    result = load_shift_requests(path)
    assert len(result.requests) == 1
    assert result.requests[0].employee.name == "Alvarez, Jordan"
    assert len(result.issues) == 1


def test_missing_employee_header_raises(tmp_path):
    path = tmp_path / "kronos.xlsx"

    def build(ws):
        ws.cell(row=1, column=1, value="Not a Kronos export")

    _save(path, build)
    with pytest.raises(KronosImportError, match="Location Schedule"):
        load_shift_requests(path)


def test_rejects_xlsm(tmp_path):
    path = tmp_path / "kronos.xlsm"
    path.write_bytes(b"not a real file")
    with pytest.raises(KronosImportError, match="Macro-enabled"):
        load_shift_requests(path)


def test_rejects_missing_file(tmp_path):
    with pytest.raises(KronosImportError, match="not found"):
        load_shift_requests(tmp_path / "does_not_exist.xlsx")


def test_prefers_a_sheet_named_report_when_multiple_sheets_exist(tmp_path):
    path = tmp_path / "kronos.xlsx"
    d1 = date(2026, 8, 1)

    def build_all(wb):
        decoy = wb.active
        decoy.title = "Cover Page"
        decoy.cell(row=1, column=1, value="Not the data")
        report = wb.create_sheet("Report")
        _write_block(report, 6, [d1], [("Alvarez, Jordan", {d1: "D12"})])

    wb = Workbook()
    build_all(wb)
    wb.save(path)

    result = load_shift_requests(path)
    assert len(result.requests) == 1
