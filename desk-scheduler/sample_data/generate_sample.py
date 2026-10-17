"""
Generates a realistic-looking fake Kronos "Location Schedule - Monthly"
export for demoing and testing the scheduler without needing a real
hospital export.

The shape here mirrors a real export (reverse engineered from an actual
Children's Colorado pull): a banner row, an "Employee" header row, a date
row below it, then one row per employee with shift codes (D12, 08c12),
avoid-markers (RDY/SDY/VDY), and occasional PTO scattered in -- not just a
simple table.

Run directly to drop a sample file next to this script:

    python -m sample_data.generate_sample
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SAMPLE_EMPLOYEES = [
    "Alvarez, Jordan",
    "Brooks, Taylor",
    "Castillo-Reyes, Morgan",
    "Dalton, Casey",
    "Ellison, Avery",
    "Foster, Quinn",
    "Grant, Riley",
    "Hayes, Jamie",
    "Irving, Drew",
]

SHIFT_CODES = ["D12", "08c12"]
SHIFT_WEIGHTS = [0.75, 0.25]  # openers are more common than the later shift

# What ends up in a "not working" day: mostly just blank, sometimes a
# soft avoid-request, rarely PTO.
NON_WORK_MARKERS = [None, "RDY", "SDY", "VDY 12:30", "VDY 00:30", "PTO Scheduled Hours 12:00"]
NON_WORK_WEIGHTS = [0.70, 0.10, 0.08, 0.05, 0.02, 0.05]


def _weekday_header(d: date) -> str:
    return d.strftime("%a")


def generate_sample_workbook(
    output_path: str | Path,
    start: date,
    end: date,
    employees: list[str] | None = None,
    shifts_per_week: int = 3,
    seed: int = 42,
) -> Path:
    """Writes a sample Kronos-grid export covering [start, end] inclusive."""
    rng = random.Random(seed)
    employees = employees or SAMPLE_EMPLOYEES
    output_path = Path(output_path)

    all_days: list[date] = []
    day = start
    while day <= end:
        all_days.append(day)
        day += timedelta(days=1)

    weeks: dict[tuple[int, int], list[date]] = {}
    for d in all_days:
        key = d.isocalendar()[:2]
        weeks.setdefault(key, []).append(d)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ws.cell(row=1, column=1, value="Location Schedule - Monthly").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Time Period :")
    ws.cell(row=2, column=4, value=f"{start.month}/{start.day:02d}/{start.year} - {end.month}/{end.day:02d}/{end.year}")
    ws.cell(row=5, column=1, value="0000100-X1 - Location: SAMPLE/DEMO/DATA/0000100-X1")

    header_row = 6
    date_row = 7
    ws.cell(row=header_row, column=1, value="Employee")
    ws.cell(row=header_row, column=3, value="Job")
    for col, d in zip(range(5, 5 + 2 * len(all_days), 2), all_days):
        ws.cell(row=header_row, column=col, value=_weekday_header(d))
        ws.cell(row=date_row, column=col, value=d.strftime("%m/%d/\n%Y"))

    col_for_date = {d: col for col, d in zip(range(5, 5 + 2 * len(all_days), 2), all_days)}

    row = date_row + 1
    for emp in employees:
        ws.cell(row=row, column=1, value=emp)
        ws.cell(row=row, column=3, value="ISS")

        working_days: set[date] = set()
        for _week_key, days_in_week in weeks.items():
            pick_count = min(shifts_per_week, len(days_in_week))
            working_days.update(rng.sample(days_in_week, pick_count))

        for d in all_days:
            col = col_for_date[d]
            if d in working_days:
                shift = rng.choices(SHIFT_CODES, weights=SHIFT_WEIGHTS, k=1)[0]
                # Occasionally show an overridden avoid-marker riding along
                # with the shift, matching what real exports look like.
                if rng.random() < 0.08:
                    marker = rng.choice(["RDY", "SDY"])
                    value = f"{marker} \n{shift}"
                else:
                    value = shift
            else:
                value = rng.choices(NON_WORK_MARKERS, weights=NON_WORK_WEIGHTS, k=1)[0]
            if value:
                ws.cell(row=row, column=col, value=value)
        row += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    today = date.today()
    period_start = today.replace(day=1)
    if period_start.month == 12:
        next_month_start = period_start.replace(year=period_start.year + 1, month=1)
    else:
        next_month_start = period_start.replace(month=period_start.month + 1)
    period_end = next_month_start - timedelta(days=1)

    out = generate_sample_workbook(
        Path(__file__).parent / "sample_kronos_export.xlsx",
        period_start,
        period_end,
    )
    print(f"Sample file written to {out}")
